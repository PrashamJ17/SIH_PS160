"""Render the saved .pptx to the PDF the SIH portal actually accepts.

The template's own instruction slide is explicit: *"You need to save the file in PDF and
upload the same on portal. No PPT, Word Doc or any other format will be supported."* So
the PDF is the submission, and the .pptx is the editable source.

There is no LibreOffice on this machine, so this draws the pages itself — reading the
**saved** presentation and placing every shape at its stored position, with vector text at
its stored size, weight, colour and alignment. Deriving the PDF from the .pptx rather than
from a second layout definition is what stops the two from drifting: there is one layout,
and the PDF is a view of it.

Pages are exactly 960 x 540 pt, which is the template's own canvas, so the result is
13.333in x 7.5in at 16:9 with no scaling anywhere.

    python -m scripts.deck.render_pdf
"""

from __future__ import annotations

import io
import sys
from functools import lru_cache
from pathlib import Path
from typing import Any, Final

import matplotlib

matplotlib.use("Agg")
import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
from matplotlib.axes import Axes
from matplotlib.backends.backend_pdf import PdfPages
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.text.text import _Paragraph
from pptx.util import Emu

REPO: Final = Path(__file__).resolve().parents[2]
DECK: Final = REPO / "docs" / "SIH2026-Highlanders-IPsec-Sentinel.pptx"
OUTPUT: Final = REPO / "docs" / "SIH2026-Highlanders-IPsec-Sentinel.pdf"
PT: Final = 12700

SYSTEM_FONTS: Final = Path("/System/Library/Fonts/Supplemental")


def _register_fonts() -> set[str]:
    """Teach matplotlib the system's own Arial and Times, and report what it has.

    macOS ships neither Calibri nor Garamond — both are Microsoft faces — so the PDF
    substitutes the nearest available: Arial for Calibri, Times New Roman for Garamond.
    The .pptx still names the originals, so a machine with Office installed renders them
    as the template intends.
    """
    from matplotlib import font_manager

    available: set[str] = set()
    for pattern in ("Arial*.ttf", "Times New Roman*.ttf"):
        for path in SYSTEM_FONTS.glob(pattern):
            try:
                font_manager.fontManager.addfont(str(path))
                available.add(font_manager.FontProperties(fname=str(path)).get_name())
            except (OSError, RuntimeError):  # pragma: no cover - font probing
                continue
    return available


AVAILABLE: Final = _register_fonts()

# PowerPoint's fonts mapped onto families matplotlib can actually resolve here. Serif
# stays serif and sans stays sans, which is what carries the template's look.
FAMILY: Final = {
    "Times New Roman": ("serif", "Times New Roman"),
    "Garamond": ("serif", "Times New Roman"),
    "Arial": ("sans-serif", "Arial"),
    "Calibri": ("sans-serif", "Arial"),
}

ALIGN: Final = {
    PP_ALIGN.LEFT: "left",
    PP_ALIGN.CENTER: "center",
    PP_ALIGN.RIGHT: "right",
    None: "left",
}

# Fallback only, for a face whose file cannot be opened. Real widths come from the font.
WIDTH_RATIO: Final = {"serif": 0.46, "sans-serif": 0.52}

FONT_FILES: Final = {
    ("Arial", False): SYSTEM_FONTS / "Arial.ttf",
    ("Arial", True): SYSTEM_FONTS / "Arial Bold.ttf",
    ("Times New Roman", False): SYSTEM_FONTS / "Times New Roman.ttf",
    ("Times New Roman", True): SYSTEM_FONTS / "Times New Roman Bold.ttf",
}


@lru_cache(maxsize=64)
def _metrics(face: str, bold: bool, size: float) -> ImageFont.FreeTypeFont | None:
    path = FONT_FILES.get((face, bold))
    if path is None or not path.is_file():
        return None
    # Measured at 100x and scaled down, so fractional point sizes stay accurate.
    return ImageFont.truetype(str(path), int(size * 100))


def advance(text: str, face: str, bold: bool, size: float, family: str) -> float:
    """How wide this text is, in points, in the face it will actually be drawn in."""
    font = _metrics(face, bold, size)
    if font is None:
        return len(text) * size * WIDTH_RATIO[family]
    return font.getlength(text) / 100.0


def pts(value: Emu | int | None) -> float:
    return (value or 0) / PT


def _rgb(colour: Any) -> str | None:
    try:
        if colour is None or colour.type is None:
            return None
        return f"#{colour.rgb}"
    except (AttributeError, TypeError, ValueError):
        return None


def _fill(shape: Any) -> str | None:
    try:
        if shape.fill.type == 1:  # solid
            return _rgb(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def _line(shape: Any) -> tuple[str | None, float]:
    try:
        colour = _rgb(shape.line.color)
        width = shape.line.width.pt if shape.line.width else 1.0
        return colour, width
    except (AttributeError, TypeError, ValueError):
        return None, 1.0


def _draw_shape(ax: Axes, shape: Any, height: float) -> None:
    """The shape's own body: fill, outline, and rounded or elliptical geometry."""
    x, y = pts(shape.left), pts(shape.top)
    w, h = pts(shape.width), pts(shape.height)
    if w <= 0 or h <= 0:
        return

    face = _fill(shape)
    edge, edge_width = _line(shape)
    if face is None and edge is None:
        return

    top = height - y - h  # matplotlib's origin is bottom-left; PowerPoint's is top-left
    name = (shape.name or "").lower()
    patch: mpatches.Patch
    if "oval" in name:
        patch = mpatches.Ellipse(
            (x + w / 2, top + h / 2),
            w,
            h,
            facecolor=face or "none",
            edgecolor=edge or "none",
            linewidth=edge_width,
            zorder=2,
        )
    elif "rounded" in name:
        patch = mpatches.FancyBboxPatch(
            (x + 6, top + 6),
            w - 12,
            h - 12,
            boxstyle="round,pad=6,rounding_size=7",
            facecolor=face or "none",
            edgecolor=edge or "none",
            linewidth=edge_width,
            zorder=2,
        )
    else:
        patch = mpatches.Rectangle(
            (x, top),
            w,
            h,
            facecolor=face or "none",
            edgecolor=edge or "none",
            linewidth=edge_width,
            zorder=2,
        )
    ax.add_patch(patch)


def _run_style(run: Any) -> dict[str, Any]:
    """One run's typography, resolved to a face this machine can draw."""
    font = run.font
    family, face = FAMILY.get(font.name or "Arial", FAMILY["Arial"])
    if face not in AVAILABLE:
        face = "DejaVu Serif" if family == "serif" else "DejaVu Sans"
    return {
        "family": family,
        "face": face,
        "size": font.size.pt if font.size else 12.0,
        "bold": bool(font.bold),
        "colour": _rgb(font.color) or "#000000",
        "underline": bool(font.underline),
    }


def _lay_out(
    paragraph: _Paragraph, usable: float
) -> tuple[list[list[tuple[str, dict[str, Any]]]], float]:
    """Break a paragraph's runs into lines that fit, keeping each run's own styling.

    Rendering a paragraph in its first run's colour was wrong on the title slide, where
    the label is black and the value navy in the same line. So runs are measured and
    placed individually, and a line breaks at a word boundary inside whichever run
    happens to cross the edge.
    """
    runs = [(run.text, _run_style(run)) for run in paragraph.runs if run.text]
    if not runs:
        return [], 12.0

    lines: list[list[tuple[str, dict[str, Any]]]] = [[]]
    used = 0.0
    for text, style in runs:
        for word in text.replace("\n", " ").split(" "):
            if not word:
                continue
            piece = word if not lines[-1] and used == 0 else " " + word
            width = advance(piece, style["face"], style["bold"], style["size"], style["family"])
            if used + width > usable and lines[-1]:
                lines.append([])
                used = 0.0
                piece = word
                width = advance(piece, style["face"], style["bold"], style["size"], style["family"])
            lines[-1].append((piece, style))
            used += width

    biggest = max(style["size"] for _, style in runs)
    return [line for line in lines if line], biggest


def _draw_text(ax: Axes, shape: Any, height: float) -> None:
    """Every paragraph, wrapped inside the shape and stacked down it."""
    if not shape.has_text_frame:
        return
    x, y = pts(shape.left), pts(shape.top)
    w, h = pts(shape.width), pts(shape.height)

    inset_l = pts(shape.text_frame.margin_left)
    inset_r = pts(shape.text_frame.margin_right)
    inset_t = pts(shape.text_frame.margin_top)
    usable = max(w - inset_l - inset_r, 20)

    blocks: list[tuple[list[tuple[str, dict[str, Any]]], float, str, bool]] = []
    for paragraph in shape.text_frame.paragraphs:
        lines, size = _lay_out(paragraph, usable)
        align = ALIGN.get(paragraph.alignment, "left")
        for index, line in enumerate(lines):
            marker = False
            if index == 0 and line and line[0][0].lstrip().startswith("\u2756"):
                # The template's section bullet is a Wingdings diamond and Arial has no
                # U+2756; it is drawn as a marker rather than left as a blank box.
                text, style = line[0]
                line[0] = (text.replace("\u2756", "").lstrip(), style)
                marker = True
            blocks.append((line, size * 1.22, align, marker))
        after = paragraph.space_after.pt if paragraph.space_after else 0
        if after:
            blocks.append(([], after, align, False))

    if not blocks:
        return

    total = sum(step for _, step, _, _ in blocks)
    anchor = shape.text_frame.vertical_anchor
    cursor = y + (h - total) / 2 if anchor is not None and int(anchor) == 3 else y + inset_t

    for line, step, align, marker in blocks:
        if not line:
            cursor += step
            continue
        widths = [
            advance(text, style["face"], style["bold"], style["size"], style["family"])
            for text, style in line
        ]
        span = sum(widths)
        if align == "center":
            pen = x + inset_l + (usable - span) / 2
        elif align == "right":
            pen = x + inset_l + usable - span
        else:
            pen = x + inset_l

        lead = line[0][1]
        baseline = height - cursor - lead["size"] * 0.94

        if marker:
            ax.add_patch(
                mpatches.RegularPolygon(
                    (pen - lead["size"] * 0.62, baseline + lead["size"] * 0.30),
                    4,
                    radius=lead["size"] * 0.34,
                    facecolor=lead["colour"],
                    edgecolor="none",
                    zorder=6,
                )
            )

        for (text, style), width in zip(line, widths, strict=True):
            ax.text(
                pen,
                baseline,
                text,
                ha="left",
                va="baseline",
                zorder=6,
                fontsize=style["size"],
                fontweight="bold" if style["bold"] else "normal",
                color=style["colour"],
                family=style["family"],
                fontname=style["face"],
            )
            if style["underline"]:
                ax.plot(
                    [pen, pen + width],
                    [baseline - style["size"] * 0.20] * 2,
                    color=style["colour"],
                    linewidth=0.9,
                    zorder=6,
                )
            pen += width
        cursor += step


def _draw_picture(ax: Axes, shape: Any, height: float) -> None:
    x, y = pts(shape.left), pts(shape.top)
    w, h = pts(shape.width), pts(shape.height)
    image = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
    ax.imshow(
        image,
        extent=(x, x + w, height - y - h, height - y),
        aspect="auto",
        zorder=4,
        interpolation="antialiased",
    )


def render(deck: Path = DECK, out: Path = OUTPUT) -> Path:
    presentation = Presentation(str(deck))
    width = pts(presentation.slide_width)
    height = pts(presentation.slide_height)

    out.parent.mkdir(parents=True, exist_ok=True)
    with PdfPages(out) as pdf:
        for index, slide in enumerate(presentation.slides, start=1):
            figure = plt.figure(figsize=(width / 72, height / 72), dpi=200)
            ax = figure.add_axes((0, 0, 1, 1))
            ax.set_xlim(0, width)
            ax.set_ylim(0, height)
            ax.axis("off")
            ax.add_patch(mpatches.Rectangle((0, 0), width, height, facecolor="white", zorder=0))

            for shape in slide.shapes:
                if shape.shape_type == 13:
                    _draw_picture(ax, shape, height)
                    continue
                _draw_shape(ax, shape, height)
                _draw_text(ax, shape, height)

            pdf.savefig(figure)
            plt.close(figure)
            print(f"  page {index}: {len(slide.shapes)} shapes")

    print(
        f"\nwrote {out.relative_to(REPO)} "
        f"({out.stat().st_size // 1024} KB, {width:.0f}x{height:.0f}pt pages)"
    )
    return out


def main() -> int:
    if not DECK.is_file():
        print(f"no deck at {DECK}; run scripts.deck.build_deck first", file=sys.stderr)
        return 1
    render()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
